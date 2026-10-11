from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter


HEADERS = ("日期", "销售员", "产品", "数量", "单价", "地区")

ROWS = (
    (date(2026, 9, 1), "张敏", "笔记本电脑", 2, 5299.00, "华东"),
    (date(2026, 9, 2), "李强", "无线鼠标", 5, 89.90, "华北"),
    (date(2026, 9, 3), "王芳", "机械键盘", 3, 399.00, "华南"),
    (date(2026, 9, 4), "陈杰", "显示器", 1, 1599.00, "华东"),
    (date(2026, 9, 5), "刘洋", "USB-C扩展坞", 4, 259.00, "西南"),
    (date(2026, 9, 6), "赵婷", "耳机", 2, 699.00, "华中"),
    (date(2026, 9, 7), "周凯", "移动硬盘", 3, 459.00, "东北"),
    (date(2026, 9, 8), "孙悦", "网络摄像头", 2, 329.00, "西北"),
    (date(2026, 9, 9), "张敏", "笔记本电脑", 2, 5299.00, "华东"),
    (date(2026, 9, 2), "李强", "无线鼠标", 5, 89.90, "华北"),
    (date(2026, 9, 11), None, "机械键盘", 1, 399.00, "华南"),
    (date(2026, 9, 12), "王芳", None, 2, 1599.00, "华东"),
    (date(2026, 9, 13), "陈杰", "移动硬盘", -2, 459.00, "西南"),
    (date(2026, 9, 14), "刘洋", "耳机", 0, 699.00, "华中"),
    (date(2026, 9, 15), "赵婷", "网络摄像头", 1, 329.00, "东北"),
)


def main() -> None:
    output_path = Path(__file__).resolve().parents[1] / "testdata" / "sales_raw_sample.xlsx"
    output_path.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "销售原始数据"
    worksheet.append(HEADERS)

    for row in ROWS:
        worksheet.append(row)

    for cell in worksheet[1]:
        cell.font = Font(bold=True)
    for row in worksheet.iter_rows(min_row=2, min_col=1, max_col=1):
        row[0].number_format = "yyyy-mm-dd"
    for row in worksheet.iter_rows(min_row=2, min_col=5, max_col=5):
        row[0].number_format = "0.00"
    for column_index, width in enumerate((14, 12, 20, 10, 12, 12), start=1):
        worksheet.column_dimensions[get_column_letter(column_index)].width = width

    workbook.save(output_path)
    print(f"已生成测试数据：{output_path}")


if __name__ == "__main__":
    main()
