import argparse
from pathlib import Path

from openpyxl import Workbook

if __package__:
    from .core import (
        DETAIL_SHEET,
        MONTH_COLUMN,
        OVERVIEW_SHEET,
        REGION_SHEET,
        REQUIRED_COLUMNS,
        REVENUE_COLUMN,
        SALESPERSON_SHEET,
        InputFileNotFoundError,
        InvalidExcelFileError,
        NoValidSalesDataError,
        generate_report,
    )
else:
    from core import (
        DETAIL_SHEET,
        MONTH_COLUMN,
        OVERVIEW_SHEET,
        REGION_SHEET,
        REQUIRED_COLUMNS,
        REVENUE_COLUMN,
        SALESPERSON_SHEET,
        InputFileNotFoundError,
        InvalidExcelFileError,
        NoValidSalesDataError,
        generate_report,
    )


def _write_empty_report(output_path: str | Path) -> None:
    """生成包含四个工作表及表头的空报表模板。"""
    destination = Path(output_path)
    destination.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    detail = workbook.active
    detail.title = DETAIL_SHEET
    detail.append([*REQUIRED_COLUMNS, REVENUE_COLUMN])

    salesperson = workbook.create_sheet(SALESPERSON_SHEET)
    salesperson.append(["销售员", REVENUE_COLUMN])

    region = workbook.create_sheet(REGION_SHEET)
    region.append(["地区", REVENUE_COLUMN])

    overview = workbook.create_sheet(OVERVIEW_SHEET)
    overview.append([MONTH_COLUMN, REVENUE_COLUMN])

    workbook.save(destination)


def main() -> int:
    parser = argparse.ArgumentParser(description="清洗销售数据并生成Excel报表")
    parser.add_argument(
        "--input",
        required=True,
        help="原始销售数据Excel文件路径",
    )
    parser.add_argument(
        "--output",
        required=True,
        help="输出报表Excel文件路径",
    )
    args = parser.parse_args()

    try:
        generate_report(args.input, args.output)
    except InputFileNotFoundError:
        print("错误：指定的输入文件不存在，请检查路径")
        return 1
    except InvalidExcelFileError:
        print("错误：输入文件不是合法的Excel文件")
        return 1
    except NoValidSalesDataError:
        try:
            _write_empty_report(args.output)
        except Exception as error:
            print(f"运行出错：{error}")
            return 1
        print("提示：原始数据清洗后没有剩余有效记录，已生成空报表模板")
    except Exception as error:
        print(f"运行出错：{error}")
        return 1

    print("报表生成完成，输出文件已保存至指定路径")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
