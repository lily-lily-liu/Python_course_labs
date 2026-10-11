from datetime import date

import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook

from sales_report.core import (
    DETAIL_SHEET,
    OVERVIEW_SHEET,
    REGION_SHEET,
    REQUIRED_COLUMNS,
    REVENUE_COLUMN,
    SALESPERSON_SHEET,
    NoValidSalesDataError,
    SalesReportError,
    generate_report,
)


def _write_sales_excel(path, rows):
    """将测试数据写入临时Excel文件。"""
    pd.DataFrame(rows, columns=REQUIRED_COLUMNS).to_excel(path, index=False)
    return path


def test_generate_report_with_valid_data_and_seller_totals(tmp_path):
    # 验证合法销售数据可生成四张工作表，且销售员销售额汇总准确。
    input_path = _write_sales_excel(
        tmp_path / "sales.xlsx",
        [
            (date(2026, 1, 2), "李雷", "键盘", 2, 100.0, "华东"),
            (date(2026, 1, 3), "李雷", "鼠标", 1, 50.0, "华东"),
            (date(2026, 1, 4), "韩梅梅", "显示器", 3, 20.0, "华南"),
        ],
    )
    output_path = tmp_path / "report.xlsx"

    generate_report(input_path, output_path)

    workbook = load_workbook(output_path, read_only=True, data_only=True)
    try:
        assert workbook.sheetnames == [
            DETAIL_SHEET,
            SALESPERSON_SHEET,
            REGION_SHEET,
            OVERVIEW_SHEET,
        ]
        summary = list(workbook[SALESPERSON_SHEET].iter_rows(values_only=True))
        assert summary[0] == ("销售员", REVENUE_COLUMN)
        assert summary[1:] == [("李雷", 250.0), ("韩梅梅", 60.0)]
    finally:
        workbook.close()


def test_generate_report_removes_rows_with_missing_key_fields(tmp_path):
    # 验证关键字段含空值的整行会被删除，其他有效记录正常保留。
    input_path = _write_sales_excel(
        tmp_path / "sales_with_blanks.xlsx",
        [
            (date(2026, 2, 1), "李雷", "键盘", 2, 100.0, "华东"),
            (date(2026, 2, 2), None, "鼠标", 1, 50.0, "华东"),
            (None, "缺少日期", "鼠标", 1, 50.0, "华东"),
            (date(2026, 2, 3), "缺少产品", None, 1, 50.0, "华东"),
            (date(2026, 2, 4), "缺少数量", "鼠标", None, 50.0, "华东"),
            (date(2026, 2, 5), "缺少单价", "鼠标", 1, None, "华东"),
            (date(2026, 2, 6), "缺少地区", "鼠标", 1, 50.0, None),
            (date(2026, 2, 3), "韩梅梅", "显示器", 3, 20.0, "华南"),
        ],
    )
    output_path = tmp_path / "report.xlsx"

    generate_report(input_path, output_path)

    workbook = load_workbook(output_path, read_only=True, data_only=True)
    try:
        detail = list(workbook[DETAIL_SHEET].iter_rows(values_only=True))
        assert len(detail) == 3
        assert all(row[1] is not None for row in detail[1:])
        assert {row[1] for row in detail[1:]} == {"李雷", "韩梅梅"}
    finally:
        workbook.close()


def test_generate_report_removes_non_positive_quantities(tmp_path):
    # 验证数量为零或负数的行均不会进入明细和销售员汇总。
    input_path = _write_sales_excel(
        tmp_path / "sales_with_invalid_quantities.xlsx",
        [
            (date(2026, 3, 1), "李雷", "键盘", 2, 100.0, "华东"),
            (date(2026, 3, 2), "负数记录", "鼠标", -1, 50.0, "华东"),
            (date(2026, 3, 3), "零数量记录", "显示器", 0, 20.0, "华南"),
        ],
    )
    output_path = tmp_path / "report.xlsx"

    generate_report(input_path, output_path)

    workbook = load_workbook(output_path, read_only=True, data_only=True)
    try:
        detail = list(workbook[DETAIL_SHEET].iter_rows(values_only=True))
        summary = list(workbook[SALESPERSON_SHEET].iter_rows(values_only=True))
        assert len(detail) == 2
        assert all(row[3] > 0 for row in detail[1:])
        assert summary[1:] == [("李雷", 200.0)]
    finally:
        workbook.close()


def test_generate_report_rejects_completely_blank_workbook(tmp_path):
    # 验证没有任何表头或数据的空白xlsx会抛出可读的自定义异常。
    input_path = tmp_path / "blank.xlsx"
    workbook = Workbook()
    workbook.save(input_path)
    output_path = tmp_path / "report.xlsx"

    with pytest.raises(SalesReportError, match="缺少必需字段"):
        generate_report(input_path, output_path)

    assert not output_path.exists()


def test_generate_report_rejects_data_with_no_valid_rows(tmp_path):
    # 验证全部记录均为空值或数量非正时会报告无有效数据且不生成残缺报表。
    input_path = _write_sales_excel(
        tmp_path / "all_dirty.xlsx",
        [
            (date(2026, 4, 1), None, "键盘", 2, 100.0, "华东"),
            (date(2026, 4, 2), "负数记录", "鼠标", -1, 50.0, "华东"),
            (date(2026, 4, 3), "零数量记录", "显示器", 0, 20.0, "华南"),
        ],
    )
    output_path = tmp_path / "report.xlsx"

    with pytest.raises(NoValidSalesDataError, match="清洗后没有有效销售数据"):
        generate_report(input_path, output_path)

    assert not output_path.exists()
